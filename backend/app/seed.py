"""Демо-данные OpenLesson.

The script loads the reference data (staff directory, schools, campus
locations, feedback criteria) and then *simulates* the platform timeline: the
real automation cycle is executed at every relevant moment while uploads,
attendance and feedback are performed "in the past". This way every status,
notification, poster and e-mail in the demo database is produced by the same
code that runs in production.

The staff directory is the university export (teachers JSON/CSV). It contains
personal data, so it lives in ``backend/data/private/`` which is excluded from
git. Without it a small fictional directory is used.

Usage:  python -m app.seed [--reset] [--teachers path/to/teachers.json]
"""

import json
import random
import re
import sys
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from pathlib import Path

import docx
from pptx import Presentation
from sqlalchemy import select

from app import clock
from app.config import BASE_DIR, settings
from app.defaults import CRITERIA
from app.db import Base, SessionLocal, engine, init_db
from app.models import (
    AcademicYear,
    Department,
    Feedback,
    FeedbackCriterion,
    FeedbackStatus,
    LessonMaterial,
    LessonStatus,
    Location,
    OpenLesson,
    School,
    Trimester,
    User,
)
from app import seed_en
from app.i18n import tr
from app.services import ai, app_settings, automation, lesson_logic, staff_directory, storage

rng = random.Random(2026)
LANG = "ru"  # demo content language: ru | en (``--lang en``)
DOMAIN = "astanait.edu.kz"
DEFAULT_DIRECTORY = BASE_DIR / "data" / "private" / "teachers.json"

# Service accounts for the demo (fictional people). Deans get a real school scope.
# (ФИО, login, должность, школа, роли, managed school)
STAFF = [
    ("Сейтказина Динара Маратовна", "d.seitkazina", "Администратор платформы (демо)", None, ["admin"], None),
    ("Жумабекова Асель Ерлановна", "a.zhumabekova", "Методист (демо)", None, ["methodist"], None),
    ("Касымов Нурлан Бакытович", "n.kassymov", "Декан (демо)", "Школа программной инженерии", ["dean"], "Школа программной инженерии"),
    ("Ибраева Гульнара Кайратовна", "g.ibrayeva", "Заместитель декана (демо)", "Школа общеобразовательных дисциплин", ["dean"], "Школа общеобразовательных дисциплин"),
    ("Омаров Ерлан Серикович", "e.omarov", "Проректор по академическим вопросам (демо)", None, ["leadership"], None),
]

# Used only when the university export is not available (e.g. a fresh clone of the repository).
FALLBACK_TEACHERS = [
    ("Ахметова Айгерим Сериковна", "Школа программной инженерии"),
    ("Нурланов Данияр Аскарулы", "Школа кибербезопасности"),
    ("Смагулова Жанар Болатовна", "Школа общеобразовательных дисциплин"),
    ("Ким Виктор Александрович", "Школа искусственного интеллекта и науки о данных"),
    ("Есенова Мадина Талгатовна", "Школа искусственного интеллекта и науки о данных"),
    ("Бекмуханов Арман Ерболович", "Школа программной инженерии"),
    ("Тулегенова Айжан Муратовна", "Школа креативной индустрии"),
    ("Абдрахманов Тимур Русланович", "Школа кибербезопасности"),
    ("Сарсенбаева Алия Нурлановна", "Школа общеобразовательных дисциплин"),
    ("Петров Алексей Игоревич", "Школа интеллектуальных систем"),
    ("Жаксылыкова Камила Ериковна", "Школа искусственного интеллекта и науки о данных"),
    ("Муратов Ерасыл Канатович", "Школа программной инженерии"),
    ("Исмаилова Сабина Рашидовна", "Школа общеобразовательных дисциплин"),
    ("Даулетов Нурсултан Маратович", "Школа интеллектуальных систем"),
    ("Оспанова Дана Сериковна", "Школа креативной индустрии"),
    ("Бейсенов Олжас Галымулы", "Школа цифрового государственного управления"),
    ("Рахимова Лаура Аскаровна", "Школа искусственного интеллекта и науки о данных"),
]


TOPICS = [
    ("Интерактивные методы обучения в программировании", "Основы программирования"),
    ("Искусственный интеллект в образовании", "Информационные технологии"),
    ("Проектирование REST API на FastAPI", "Веб-разработка"),
    ("Нормализация реляционных баз данных", "Базы данных"),
    ("Атаки на веб-приложения и методы защиты (OWASP Top 10)", "Информационная безопасность"),
    ("Линейная регрессия и метод наименьших квадратов", "Машинное обучение"),
    ("Сторителлинг в цифровых медиа", "Цифровая журналистика"),
    ("Academic Writing: структура научной статьи", "Academic English"),
    ("Графы и алгоритмы поиска кратчайшего пути", "Дискретная математика"),
    ("Контейнеризация приложений с Docker", "DevOps"),
    ("Сетевые протоколы TCP/IP на практике", "Компьютерные сети"),
    ("Agile и Scrum: роли и артефакты", "Управление IT-проектами"),
    ("История Казахстана в цифровую эпоху", "Современная история Казахстана"),
    ("Численные методы решения уравнений", "Вычислительная математика"),
    ("UX-исследования и прототипирование в Figma", "Дизайн интерфейсов"),
    ("Микроконтроллеры и IoT-устройства", "Компьютерная архитектура"),
    ("Нейронные сети: от перцептрона к CNN", "Глубокое обучение"),
    ("Критическое мышление и аргументация", "Философия"),
    ("Тестирование ПО: unit и интеграционные тесты", "Качество программного обеспечения"),
    ("Kazakh for IT: профессиональная лексика", "Казахский язык"),
    ("Вероятностные модели в анализе данных", "Теория вероятностей"),
    ("Криптография с открытым ключом", "Криптография"),
    ("Медиаграмотность и фактчекинг", "Медиа и коммуникации"),
    ("Облачные вычисления: модели IaaS, PaaS, SaaS", "Облачные технологии"),
    ("Паттерны проектирования GoF", "Объектно-ориентированное проектирование"),
    ("Публичные выступления на английском языке", "Professional English"),
    ("Big Data: экосистема Hadoop и Spark", "Обработка больших данных"),
    ("Экономика цифровых платформ", "Цифровая экономика"),
    ("Компьютерное зрение с OpenCV", "Компьютерное зрение"),
    ("Мобильная разработка на Flutter", "Мобильная разработка"),
    ("Этика искусственного интеллекта", "ИИ и общество"),
    ("Матрицы и линейные преобразования", "Линейная алгебра"),
]

STRENGTHS = [
    "Чёткая структура занятия, каждый этап логично вытекал из предыдущего.",
    "Отличное использование Mentimeter для проверки понимания в режиме реального времени.",
    "Студенты активно работали в группах, атмосфера была доброжелательной и рабочей.",
    "Практические примеры из индустрии сделали материал понятным и актуальным.",
    "Преподаватель уверенно владеет материалом и умело управляет дискуссией.",
    "Хорошо продуманная рефлексия в конце занятия — студенты сами сформулировали выводы.",
    "Интересный кейс, который студенты решали в командах, вызвал живое обсуждение.",
    "Наглядная презентация, удачные визуализации сложных понятий.",
]
IMPROVEMENTS = [
    "Можно оставить больше времени на вопросы студентов в конце занятия.",
    "Темп второй части был немного высоким — часть студентов не успевала конспектировать.",
    "Стоило бы чаще вовлекать студентов с задних рядов.",
    "Рекомендую добавить короткий квиз в начале для актуализации знаний.",
    "Слайды местами перегружены текстом — лучше сократить и добавить схемы.",
    "Домашнее задание стоит связать с практической частью занятия.",
]
COMMENTS = [
    "Спасибо за открытое занятие, возьму несколько приёмов в свою практику.",
    "Полезный опыт для коллег кафедры.",
    "",
    "Хотелось бы увидеть продолжение темы на следующем открытом занятии.",
    "Занятие соответствует современным требованиям к активному обучению.",
]


def email(login: str) -> str:
    return f"{login}@{DOMAIN}"


# ------------------------------------------------------------------ reference data


def load_directory(path: Path | None) -> list[staff_directory.Person]:
    if path and path.exists():
        people = staff_directory.parse(path.read_bytes(), path.name)
        print(f"Справочник ППС: {path} — {len(people)} преподавателей")
        return people
    print("Файл справочника ППС не найден — используется вымышленный список преподавателей")
    if LANG == "en":
        return [
            staff_directory.Person(external_id=f"demo-{i}", full_name=name, school_ru=school, school_kz=None, school_en=school, school_id=None)
            for i, (name, school) in enumerate(seed_en.FALLBACK_TEACHERS, 1)
        ]
    return [
        staff_directory.Person(external_id=f"demo-{i}", full_name=name, school_ru=school, school_kz=None, school_en=None, school_id=None)
        for i, (name, school) in enumerate(FALLBACK_TEACHERS, 1)
    ]


def seed_reference(db, directory: Path | None) -> None:
    result = staff_directory.import_people(db, load_directory(directory))
    print(f"  создано: {result['created']}, школ: {result['schools']}, без школы: {result['without_school']}")
    schools = {s.name: s for s in db.scalars(select(School)).all()}
    depts = {d.name: d for d in db.scalars(select(Department)).all()}
    staff = seed_en.STAFF if LANG == "en" else STAFF
    dean_schools = (seed_en.SCHOOL_SE, seed_en.SCHOOL_GE) if LANG == "en" else ("Школа программной инженерии", "Школа общеобразовательных дисциплин")
    for name in dean_schools:
        if name not in schools:  # make sure the demo deans have a scope even with a partial export
            schools[name] = School(name=name, name_en=name if LANG == "en" else None)
            db.add(schools[name])
            db.flush()
            depts[name] = Department(name=name, school_id=schools[name].id)
            db.add(depts[name])
            db.flush()
    for full_name, login, position, dept, roles, managed in staff:
        db.add(
            User(
                email=email(login),
                external_id=f"service-{login}",
                full_name=full_name,
                position=position,
                department_id=depts[dept].id if dept else None,
                managed_school_id=schools[managed].id if managed else None,
                roles=roles,
                is_faculty=False,
                language=LANG,
                created_at=clock.now(),
            )
        )
    rooms = json.loads((BASE_DIR / "app" / "data" / "rooms.json").read_text(encoding="utf-8"))
    for r in rooms:
        db.add(Location(**r))
    for i, (title, description, title_en, description_en) in enumerate(CRITERIA, 1):
        db.add(FeedbackCriterion(title=title, description=description, title_en=title_en, description_en=description_en, order=i))
    db.commit()


def user_key(u: User) -> str:
    return u.email.split("@")[0] if u.external_id and u.external_id.startswith("service-") else f"t{u.external_id}"


CYRILLIC_FULL = re.compile(r"^[А-ЯЁӘІҢҒҮҰҚӨҺа-яёәіңғүұқөһ-]+( [А-ЯЁӘІҢҒҮҰҚӨҺа-яёәіңғүұқөһ-]+){2}$")


def lesson_teachers(users: dict[str, User], count: int) -> list[str]:
    """Deterministic sample of teachers for demo lessons: assigned to a school, full Cyrillic name."""
    pool = sorted(
        (k for k, u in users.items() if "teacher" in u.roles and u.department_id and CYRILLIC_FULL.match(u.full_name)),
        key=lambda k: users[k].external_id or "",
    )
    if len(pool) < count:
        pool += sorted(k for k, u in users.items() if "teacher" in u.roles and k not in pool)
    picked = random.Random(424).sample(pool, min(count, len(pool)))
    while len(picked) < count:  # tiny fallback directory: reuse teachers
        picked += picked[: count - len(picked)]
    return picked


# ------------------------------------------------------------------ documents


def make_plan(path: Path, lesson: OpenLesson, quality: str) -> None:
    if LANG == "en":
        return make_plan_en(path, lesson, quality)
    d = docx.Document()
    d.add_heading(f"План открытого занятия: {lesson.topic}", level=1)
    d.add_paragraph(f"Дисциплина: {lesson.discipline}")
    d.add_paragraph(f"Преподаватель: {lesson.teacher.full_name}, {lesson.teacher_position}")
    d.add_paragraph(f"Продолжительность: {int((lesson.end_dt - lesson.start_dt).seconds / 60)} минут")
    d.add_heading("Цель занятия", level=2)
    d.add_paragraph(f"Сформировать у студентов понимание темы «{lesson.topic}» и навыки применения на практике.")
    if quality != "weak":
        d.add_heading("Задачи", level=2)
        for t in ("актуализировать базовые знания", "объяснить ключевые понятия", "закрепить материал на практическом задании"):
            d.add_paragraph(t, style="List Bullet")
        d.add_heading("Ожидаемые результаты обучения", level=2)
        d.add_paragraph("По итогам занятия студенты будут знать основные понятия темы и уметь применять их при решении задач; "
                        "результат проверяется мини-квизом.")
    d.add_heading("Ход занятия", level=2)
    stages = [("Организационный момент, постановка цели", 5), ("Актуализация знаний: квиз в Mentimeter", 10),
              ("Объяснение нового материала с презентацией", 15), ("Работа в группах над кейсом, обсуждение", 15),
              ("Рефлексия и подведение итогов", 5)]
    if quality == "weak":
        stages = [("Объяснение нового материала", 35), ("Ответы на вопросы", 15)]
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "Этап", "Время, мин"
    for stage, minutes in stages:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = stage, str(minutes)
    if quality == "strong":
        d.add_heading("Интерактивные методы и цифровые инструменты", level=2)
        d.add_paragraph("Дискуссия, peer review, кейс-метод; цифровые инструменты: Mentimeter, Miro, Moodle (LMS).")
    d.save(path)


def make_plan_en(path: Path, lesson: OpenLesson, quality: str) -> None:
    P = seed_en.PLAN
    d = docx.Document()
    d.add_heading(P["heading"].format(topic=lesson.topic), level=1)
    d.add_paragraph(P["discipline"].format(discipline=lesson.discipline))
    d.add_paragraph(P["teacher"].format(teacher=lesson.teacher.full_name, position=lesson.teacher_position))
    d.add_paragraph(P["duration"].format(minutes=int((lesson.end_dt - lesson.start_dt).seconds / 60)))
    d.add_heading(P["goal_h"], level=2)
    d.add_paragraph(P["goal"].format(topic=lesson.topic))
    if quality != "weak":
        d.add_heading(P["tasks_h"], level=2)
        for t in P["tasks"]:
            d.add_paragraph(t, style="List Bullet")
        d.add_heading(P["outcomes_h"], level=2)
        d.add_paragraph(P["outcomes"])
    d.add_heading(P["flow_h"], level=2)
    stages = P["stages_weak"] if quality == "weak" else P["stages"]
    table = d.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = P["stage"], P["minutes"]
    for stage, minutes in stages:
        row = table.add_row()
        row.cells[0].text, row.cells[1].text = stage, str(minutes)
    if quality == "strong":
        d.add_heading(P["methods_h"], level=2)
        d.add_paragraph(P["methods"])
    d.save(path)


def make_presentation(path: Path, lesson: OpenLesson) -> None:
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[0])
    s.shapes.title.text = lesson.topic
    s.placeholders[1].text = f"{lesson.teacher.full_name}\n{lesson.discipline}"
    titles = seed_en.PLAN["slides"] if LANG == "en" else ("Цели занятия", "Ключевые понятия", "Практическое задание", "Рефлексия")
    bullets = seed_en.PLAN["bullets"] if LANG == "en" else "• Пункт 1\n• Пункт 2\n• Пункт 3"
    for title in titles:
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = title
        slide.placeholders[1].text = bullets
    prs.save(path)


def upload(db, lesson: OpenLesson, kind: str) -> None:
    """Simulates the teacher's upload at the current (simulated) time."""
    if lesson.material(kind) or lesson.status not in lesson_logic.UPLOAD_ALLOWED:
        return
    folder = storage.subdir("materials")
    if kind == "plan":
        stored = f"seed_{lesson.id}_plan.docx"
        make_plan(folder / stored, lesson, rng.choice(["strong", "strong", "medium", "weak"]))
        prefix = "Open_lesson_plan" if LANG == "en" else "План_открытого_занятия"
        original = f"{prefix}_{lesson.teacher.full_name.split()[0]}.docx"
    else:
        stored = f"seed_{lesson.id}_presentation.pptx"
        make_presentation(folder / stored, lesson)
        original = f"{'Slides' if LANG == 'en' else 'Презентация'}_{lesson.teacher.full_name.split()[0]}.pptx"
    now = clock.now()
    deadline = lesson.effective_deadline
    lesson.materials.append(
        LessonMaterial(
            kind=kind,
            original_name=original,
            stored_name=stored,
            content_type=None,
            size=(folder / stored).stat().st_size,
            uploaded_at=now,
            uploaded_by_id=lesson.teacher_id,
            on_time=deadline is None or now <= deadline,
        )
    )
    lesson_logic.sync_materials_status(lesson)
    title = lesson_logic.MATERIAL_TITLES[kind]
    lesson_logic.log_event(db, lesson, "material_uploaded", f"{title}: файл «{original}» загружен", lesson.teacher)
    db.commit()
    if kind == "plan":
        ai.analyze_plan(db, lesson)


# ------------------------------------------------------------------ timeline simulation


def _default_uploads() -> dict:
    first = rng.choice([13, 12, 11, 9, 8])
    return {"presentation": first, "plan": max(first - rng.choice([0, 1]), 8)}


@dataclass
class Spec:
    teacher: str
    days: int  # offset of the lesson date from "today"
    start: str = "10:00"
    minutes: int = 50
    room: str | None = None
    uploads: dict = field(default_factory=lambda: _default_uploads())  # days before lesson; None = never
    observers: int = 4
    qr_share: float = 0.0  # optional QR check-in (spec 7.2) is off in the demo
    fb_share: float = 0.8  # share of observers who submit feedback
    confirm_after_h: float | None = 3  # organizer confirms the attendance list N hours after the end
    feedback_authors: list = field(default_factory=list)  # user keys forced to be observers (pending feedback)


def _rooms(db) -> list[Location]:
    return list(db.scalars(select(Location).where(Location.kind.in_(["lecture", "practice", "computer", "classroom"]))).all())


def create_lessons(db, trimester: Trimester, specs: list[Spec], today: date, topics: list, users: dict) -> list[tuple[OpenLesson, Spec]]:
    rooms = _rooms(db)
    result = []
    for spec in specs:
        teacher = users[spec.teacher]
        topic, discipline = topics.pop(0) if topics else rng.choice(seed_en.TOPICS if LANG == "en" else TOPICS)
        start = datetime.strptime(spec.start, "%H:%M").time()
        end = (datetime.combine(date.today(), start) + timedelta(minutes=spec.minutes)).time()
        loc = db.scalar(select(Location).where(Location.code == spec.room)) if spec.room else rng.choice(rooms)
        d = today + timedelta(days=spec.days)
        if abs(spec.days) > 3 and d.weekday() >= 5:  # weekend -> nearest weekday
            d += timedelta(days=-1 if d.weekday() == 5 else 1)
        lesson = OpenLesson(
            trimester_id=trimester.id,
            teacher_id=teacher.id,
            teacher_position=teacher.position,
            department_id=teacher.department_id,
            topic=topic,
            discipline=discipline,
            date=d,
            start_time=start,
            end_time=end,
            location_id=loc.id,
            location_text=loc.code,
            status=LessonStatus.draft.value,
            materials_deadline=lesson_logic.compute_deadline(db, d),
        )
        db.add(lesson)
        db.flush()
        lesson_logic.log_event(db, lesson, "created", "Занятие загружено из утверждённого графика (Excel)", users["a.zhumabekova"])
        result.append((lesson, spec))
    db.commit()
    return result


def simulate(db, items: list[tuple[OpenLesson, Spec]], users: dict, until: datetime) -> None:
    cfg = app_settings.get_all(db)
    methodist = users["a.zhumabekova"]
    faculty = [u for u in users.values() if u.is_faculty]
    events: list[tuple[datetime, int, callable]] = []

    def at(t: datetime, fn, order: int = 1):
        if t <= until:
            events.append((t, order, fn))

    for lesson, spec in items:
        start, end = lesson.start_dt, lesson.end_dt
        # automation checkpoints
        for days in (cfg["pre_reminder_days"], cfg["reminder_days"], cfg["poster_days"]):
            at(start - timedelta(days=days) + timedelta(minutes=5), None, 2)
        at(lesson.materials_deadline + timedelta(minutes=1), None, 2)
        at(end + timedelta(minutes=1), None, 2)
        # teacher uploads
        for kind, days_before in spec.uploads.items():
            if days_before is not None:
                t = datetime.combine(lesson.date - timedelta(days=days_before), time(rng.randint(9, 17), rng.choice([5, 20, 40])))
                at(t, (lambda l=lesson, k=kind: upload(db, l, k)))
        # attendance (registered during/after the lesson)
        pool = [u for u in faculty if u.id != lesson.teacher_id]
        forced = [users[login] for login in spec.feedback_authors]
        observers = forced + rng.sample([u for u in pool if u not in forced], max(spec.observers - len(forced), 0))

        def register(l=lesson, s=spec, obs=observers):
            """The methodologist enters the attendees from the paper attendance sheet."""
            if l.status != LessonStatus.completed.value:
                return
            for u in obs:
                lesson_logic.register_attendance(db, l, u, source="manual", actor=methodist, confirmed=True)
            names = ", ".join(u.full_name for u in obs)
            lesson_logic.log_event(db, l, "attendance_added", lambda: tr(
                f"Внесены посетители по явочному листу: {names}", f"Attendees entered from the attendance sheet: {names}",
                f"Қатысу парағы бойынша қатысушылар енгізілді: {names}",
            ), methodist)
            db.commit()

        if spec.confirm_after_h is not None:
            confirm_t = end + timedelta(hours=spec.confirm_after_h)
            at(confirm_t - timedelta(minutes=20), register, 0)

            def confirm(l=lesson):
                if l.status == LessonStatus.completed.value:
                    lesson_logic.confirm_attendance(db, l, methodist, None)

            at(confirm_t, confirm)
            submitters = [u for u in observers if u not in forced]
            submitters = rng.sample(submitters, round(len(submitters) * spec.fb_share))
            for u in submitters:
                t = confirm_t + timedelta(hours=rng.uniform(2, 60))
                at(t, (lambda l=lesson, a=u: submit_feedback(db, l, a)))
            due = end + timedelta(days=int(cfg["feedback_days"]))
            for d_ in (1, 2):
                at(confirm_t + timedelta(days=d_, minutes=5), None, 2)
            at(max(due, confirm_t + timedelta(hours=24)) + timedelta(minutes=2), None, 2)

    events.sort(key=lambda e: (e[0], e[1]))
    for t, _, fn in events:
        clock.set_offset(int((t - clock.real_now()).total_seconds() // 60))
        if fn is not None:
            fn()
        automation.run_cycle(db)
    clock.set_offset(0)


def submit_feedback(db, lesson: OpenLesson, author: User) -> None:
    fb = db.scalar(select(Feedback).where(Feedback.lesson_id == lesson.id, Feedback.author_id == author.id))
    if fb is None or fb.status not in ("pending", "draft") or clock.now() >= fb.due_at:
        return
    # demo feedback is attributed to real staff names, so it is kept constructive (scores 3–5)
    base = rng.choice([5, 5, 4, 4, 4])
    criteria = db.scalars(select(FeedbackCriterion)).all()
    fb.scores = {str(c.id): max(3, min(5, base + rng.choice([0, 0, 0, 1, -1]))) for c in criteria}
    strengths, improvements, comments = (
        (seed_en.STRENGTHS, seed_en.IMPROVEMENTS, seed_en.COMMENTS) if LANG == "en" else (STRENGTHS, IMPROVEMENTS, COMMENTS)
    )
    fb.strengths = " ".join(rng.sample(strengths, 2))
    fb.improvements = rng.choice(improvements)
    fb.comment = rng.choice(comments) or None
    fb.status = FeedbackStatus.submitted.value
    fb.submitted_at = clock.now()
    name = author.full_name
    lesson_logic.log_event(db, lesson, "feedback_submitted", lambda: tr(
        f"Получен отзыв от {name}", f"Review received from {name}", f"{name} пікір қалдырды"
    ), author)
    db.commit()


# ------------------------------------------------------------------ main


def run(reset: bool, directory: Path | None, lang: str = "ru") -> None:
    global LANG
    LANG = lang if lang in ("ru", "en") else "ru"
    if reset:
        Base.metadata.drop_all(bind=engine)
        for sub in ("materials", "posters", "outbox"):
            for f in (settings.storage_dir / sub).glob("*"):
                f.unlink()
    init_db()
    settings.ai_enabled = False  # demo data uses the local analyser (no API costs)
    with SessionLocal() as db:
        if db.scalar(select(User).limit(1)):
            print("База уже содержит данные. Запустите с --reset для пересоздания.")
            return
        clock.set_offset(0)
        today = clock.now().date()
        seed_reference(db, directory)
        if LANG == "en":  # English demo: e-mails, AI texts and posters in English
            for u in db.scalars(select(User)).all():
                u.language = "en"
            app_settings.set_many(db, {"poster_language": "en"})
            db.commit()
        users = {user_key(u): u for u in db.scalars(select(User)).all()}
        topics = list(seed_en.TOPICS if LANG == "en" else TOPICS)
        rng.shuffle(topics)
        topics.insert(0, seed_en.MOCKUP_TOPIC if LANG == "en" else ("Интерактивные методы обучения", "Педагогика"))  # mock-up lesson from the ТЗ

        # T[0] demo teacher, T[1] demo observer, T[2:14] previous year, T[14:30] current trimester, T[30:35] drafts
        T = lesson_teachers(users, 35)
        demo_teacher = T[0]
        observer_demo = T[1]  # a teacher who has a Feedback form to fill in right now

        # ----- previous academic year: everything finished (analytics history)
        prev = AcademicYear(name=f"{today.year - 1}–{today.year}", start_date=date(today.year - 1, 9, 1), end_date=date(today.year, 6, 30))
        prev.trimesters = [Trimester(number=n) for n in (1, 2, 3)]
        cur = AcademicYear(name=f"{today.year}–{today.year + 1}", start_date=date(today.year, 9, 1), end_date=date(today.year + 1, 6, 30), is_current=True)
        cur.trimesters = [Trimester(number=n) for n in (1, 2, 3)]
        db.add_all([prev, cur])
        db.commit()

        prev_specs = []
        for i, tri in enumerate(prev.trimesters):
            base = (date(today.year - 1, 9, 1) + timedelta(days=95 * i) - today).days
            tri.start_date = today + timedelta(days=base)
            tri.end_date = tri.start_date + timedelta(days=84)
            for j in range(4):
                prev_specs.append((tri, Spec(teacher=T[2 + i * 4 + j], days=base + 14 + j * 16,
                                             start=rng.choice(["09:00", "10:00", "11:00", "13:00", "14:00", "15:00"]),
                                             observers=rng.randint(3, 6))))
        # one annulled lesson in the previous year; the demo teacher already has a finished lesson with feedback
        prev_specs[5][1].uploads = {"presentation": 9, "plan": None}
        prev_specs[-1] = (prev.trimesters[2], Spec(teacher=demo_teacher, days=prev_specs[-1][1].days, start="10:00", observers=5, fb_share=1.0))

        # ----- current trimester I (relative to today)
        t1, t2 = cur.trimesters[0], cur.trimesters[1]
        t1.start_date, t1.end_date = today - timedelta(days=45), today + timedelta(days=40)
        t2.start_date, t2.end_date = t1.end_date + timedelta(days=7), t1.end_date + timedelta(days=95)
        cur.trimesters[2].start_date, cur.trimesters[2].end_date = t2.end_date + timedelta(days=7), t2.end_date + timedelta(days=95)
        c = iter(T[14:])
        cur_specs = [
            Spec(next(c), -30, "11:00", observers=5),
            Spec(next(c), -24, "14:00", 80, observers=4, fb_share=0.75),
            Spec(next(c), -19, "10:00", uploads={"presentation": 9, "plan": None}),
            Spec(next(c), -13, "13:00", observers=6),
            Spec(next(c), -8, "09:00", 80, observers=5),
            Spec(next(c), -2, "15:00", observers=4, fb_share=0.5, feedback_authors=[demo_teacher, observer_demo]),
            Spec(next(c), -1, "11:00", observers=5, confirm_after_h=None),  # held; attendees not entered yet
            Spec(next(c), 1, "10:00", room="C1.2.243K"),
            Spec(next(c), 2, "14:00"),
            Spec(next(c), 5, "11:00"),
            Spec(next(c), 6, "09:00", uploads={"presentation": None, "plan": None}),
            Spec(next(c), 8, "13:00", 80, uploads={"presentation": 11, "plan": 10}),
            Spec(demo_teacher, 9, "10:00", room="C1.1.255P", uploads={"presentation": 10, "plan": None}),
            Spec(next(c), 12, "15:00", uploads={"presentation": None, "plan": None}),
            Spec(next(c), 20, "10:00"),
            Spec(next(c), 27, "14:00"),
            Spec(next(c), 33, "11:00"),
        ]

        def at_time(moment: datetime) -> None:
            clock.set_offset(int((moment - clock.real_now()).total_seconds() // 60))

        # schedules are imported and approved at their historical dates
        prev_items, cur_items = [], []
        for tri in prev.trimesters + [t1]:
            approve_at = datetime.combine(tri.start_date - timedelta(days=10), time(10, 0))
            at_time(approve_at - timedelta(days=2))
            specs = cur_specs if tri is t1 else [s for t, s in prev_specs if t is tri]
            items = create_lessons(db, tri, specs, today, topics, users)
            (cur_items if tri is t1 else prev_items).extend(items)
            at_time(approve_at)
            lesson_logic.approve_trimester(db, tri, users["a.zhumabekova"])
        clock.set_offset(0)
        print("Моделирование временной шкалы…")
        simulate(db, prev_items + cur_items, users, until=clock.now())

        # ----- trimester II: draft schedule loaded but not yet approved
        draft_specs = [Spec(t, 55 + i * 9, rng.choice(["10:00", "13:00", "15:00"])) for i, t in enumerate(T[30:35])]
        create_lessons(db, t2, draft_specs, today, topics, users)

        automation.run_cycle(db)
        lessons = db.scalars(select(OpenLesson)).all()
        from collections import Counter

        print("Готово. Статусы занятий:", dict(Counter(l.status for l in lessons)))
        print("Демо-вход на странице входа. Преподаватель с занятием и формой Feedback:", users[demo_teacher].full_name)


if __name__ == "__main__":
    args = sys.argv[1:]
    path = Path(args[args.index("--teachers") + 1]) if "--teachers" in args else DEFAULT_DIRECTORY
    lang = args[args.index("--lang") + 1] if "--lang" in args else "ru"
    run(reset="--reset" in args, directory=path, lang=lang)
