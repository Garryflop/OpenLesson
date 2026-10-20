"""Excel import of the staff directory and the trimester schedule (+ templates).

Templates are generated in the request language; the parser accepts the Russian, the English and
the Kazakh column names (and a few common synonyms), so a file filled in any language imports.
Header matching is tolerant: case, extra spaces, «ё», and Kazakh letters typed with their Russian
look-alikes («Окытушынын аты-жони» = «Оқытушының аты-жөні») are ignored.
"""

import io
import re
from datetime import date, datetime, time

from fastapi import HTTPException
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.i18n import tr
from app.models import Department, LessonStatus, Location, OpenLesson, School, Trimester, TrimesterStatus, User
from app.services import lesson_logic
from app.utils import normalize_name

Fields = dict[str, tuple[str, str, str, tuple[str, ...]]]

# field key → (Russian header, English header, Kazakh header, extra accepted aliases in any language)
SCHEDULE_FIELDS: Fields = {
    "year": ("Учебный год", "Academic year", "Оқу жылы", ("Academic Year", "Year")),
    "trimester": ("Триместр", "Trimester", "Триместр", ("Term",)),
    "teacher": ("ФИО преподавателя", "Teacher full name", "Оқытушының аты-жөні",
                ("Teacher", "Full name", "ФИО", "Оқытушы", "Аты-жөні", "Т.А.Ә.", "ТАӘ")),
    "position": ("Должность", "Position", "Лауазымы", ("Лауазым",)),
    "department": ("Кафедра", "Department", "Кафедра", ()),
    "topic": ("Тема открытого занятия", "Open lesson topic", "Ашық сабақтың тақырыбы",
              ("Topic", "Тема", "Тақырыбы", "Тақырып", "Сабақ тақырыбы")),
    "discipline": ("Дисциплина", "Discipline", "Пән", ("Course", "Пәні", "Пән атауы")),
    "date": ("Дата", "Date", "Күні", ("Күн", "Өтетін күні")),
    "time": ("Время", "Time", "Уақыты", ("Уақыт", "Өтетін уақыты")),
    "place": ("Место", "Location", "Орны", ("Place", "Venue", "Room", "Аудитория", "Орын", "Өтетін орны")),
    "email": ("Email", "Email", "Email", ("E-mail", "Электрондық пошта", "Эл. пошта")),
}
STAFF_FIELDS: Fields = {
    "full_name": ("ФИО", "Full name", "Аты-жөні", ("Name", "Т.А.Ә.", "ТАӘ", "Оқытушының аты-жөні")),
    "email": ("Email", "Email", "Email", ("E-mail", "Электрондық пошта", "Эл. пошта")),
    "position": ("Должность", "Position", "Лауазымы", ("Лауазым",)),
    "department": ("Кафедра", "Department", "Кафедра", ()),
    "school": ("Школа", "School", "Мектеп", ("Мектебі",)),
    "roles": ("Роли", "Roles", "Рөлдері", ("Рөлі", "Рөлдер")),
}
OPTIONAL = {"email", "roles", "school"}

# Russian headers (kept for backward compatibility)
SCHEDULE_COLUMNS = [ru for ru, *_ in SCHEDULE_FIELDS.values()]
STAFF_COLUMNS = [ru for ru, *_ in STAFF_FIELDS.values()]

HEADER_FILL = PatternFill("solid", fgColor="0B2A5B")


def _headers(fields: Fields) -> list[str]:
    """Column names in the current request language."""
    return [tr(ru, en, kk) for ru, en, kk, _ in fields.values()]


def _workbook(columns: list[str], example: list, title: str) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = title
    ws.append(columns)
    ws.append(example)
    for i, _ in enumerate(columns, 1):
        cell = ws.cell(row=1, column=i)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = 22
    ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def schedule_template() -> bytes:
    return _workbook(
        _headers(SCHEDULE_FIELDS),
        ["2026–2027", "I", tr("Ахметова Айгерим Сериковна", "Akhmetova Aigerim Serikovna", "Ахметова Айгерим Сериковна"),
         tr("Старший преподаватель", "Senior Lecturer", "Аға оқытушы"),
         tr("Программная инженерия", "Software Engineering", "Бағдарламалық инженерия"),
         tr("Интерактивные методы обучения", "Interactive teaching methods", "Интерактивті оқыту әдістері"),
         tr("Педагогика", "Pedagogy", "Педагогика"),
         "15.10.2026", "10:00–10:50", "C1.2.243K", "a.akhmetova@astanait.edu.kz"],
        tr("График", "Schedule", "Кесте"),
    )


def staff_template() -> bytes:
    return _workbook(
        _headers(STAFF_FIELDS),
        [tr("Ахметова Айгерим Сериковна", "Akhmetova Aigerim Serikovna", "Ахметова Айгерим Сериковна"),
         "a.akhmetova@astanait.edu.kz",
         tr("Старший преподаватель", "Senior Lecturer", "Аға оқытушы"),
         tr("Программная инженерия", "Software Engineering", "Бағдарламалық инженерия"),
         tr("Школа информационных технологий и инженерии", "School of Information Technology and Engineering",
            "Ақпараттық технологиялар және инженерия мектебі"), "teacher"],
        tr("ППС", "Faculty", "ПОҚ"),
    )


# Kazakh letters → Russian look-alikes, so a header typed without a Kazakh keyboard still matches
_KK_FOLD = str.maketrans("әғқңөұүһі", "агкноуухи")


def _norm_header(value) -> str:
    text = str(value or "").replace("\xa0", " ").replace("ё", "е").replace("Ё", "Е")
    return " ".join(text.split()).lower().translate(_KK_FOLD)


def _rows(content: bytes, fields: Fields) -> list[tuple[int, dict]]:
    """Rows of the first sheet as {field key: value}; headers may be Russian, English or Kazakh."""
    try:
        wb = load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(400, tr(f"Не удалось прочитать Excel-файл: {exc}", f"Could not read the Excel file: {exc}",
                                    f"Excel файлын оқу мүмкін болмады: {exc}")) from exc
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise HTTPException(400, tr("Файл пуст", "The file is empty", "Файл бос"))
    header = [_norm_header(c) for c in rows[0]]
    index = {}
    for key, (ru, en, kk, aliases) in fields.items():
        for name in (ru, en, kk, *aliases):
            if _norm_header(name) in header:
                index[key] = header.index(_norm_header(name))
                break
    missing = [tr(ru, en, kk) for key, (ru, en, kk, _) in fields.items() if key not in OPTIONAL and key not in index]
    if missing:
        names = ", ".join(missing)
        raise HTTPException(400, tr(f"Не найдены столбцы: {names}", f"Columns not found: {names}",
                                    f"Бағандар табылмады: {names}"))
    result = []
    for n, row in enumerate(rows[1:], start=2):
        if not any(v not in (None, "") for v in row):
            continue
        result.append((n, {key: (row[i] if i < len(row) else None) for key, i in index.items()}))
    return result


def _str(v) -> str:
    return str(v).strip() if v is not None else ""


def _department(db: Session, name: str, school_name: str = "") -> Department | None:
    if not name:
        return None
    dept = db.scalar(select(Department).where(func.lower(Department.name) == name.lower()))
    if dept is None:
        school = None
        if school_name:
            school = db.scalar(select(School).where(func.lower(School.name) == school_name.lower()))
            if school is None:
                school = School(name=school_name)
                db.add(school)
                db.flush()
        dept = Department(name=name, school_id=school.id if school else None)
        db.add(dept)
        db.flush()
    return dept


def import_staff(db: Session, content: bytes) -> dict:
    created = updated = 0
    errors = []
    for n, row in _rows(content, STAFF_FIELDS):
        name, email = _str(row.get("full_name")), _str(row.get("email")).lower()
        if not name or "@" not in email:
            errors.append({"row": n, "error": tr("Не указаны ФИО или корректный Email", "Full name or a valid e-mail is missing",
                                                 "Аты-жөні немесе дұрыс Email көрсетілмеген")})
            continue
        roles = [r.strip() for r in re.split(r"[,;\s]+", _str(row.get("roles"))) if r.strip()] or ["teacher"]
        dept = _department(db, _str(row.get("department")), _str(row.get("school")))
        user = db.scalar(select(User).where(func.lower(User.email) == email))
        if user is None:
            user = User(email=email, full_name=name, roles=sorted(set(roles)))
            db.add(user)
            created += 1
        else:
            updated += 1
        user.full_name = name
        user.position = _str(row.get("position")) or user.position
        if dept:
            user.department_id = dept.id
    db.commit()
    return {"created": created, "updated": updated, "errors": errors}


def _parse_date(v) -> date:
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    s = _str(v)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y", "%d.%m.%y"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            pass
    raise ValueError(tr(f"Неверная дата «{s}» (ожидается ДД.ММ.ГГГГ)", f"Invalid date “{s}” (expected DD.MM.YYYY)",
                        f"Күні қате: «{s}» (КК.АА.ЖЖЖЖ пішімі күтіледі)"))


def _parse_time_range(v) -> tuple[time, time]:
    m = re.match(r"^\s*(\d{1,2})[:.](\d{2})\s*[-–—]\s*(\d{1,2})[:.](\d{2})\s*$", _str(v))
    if not m:
        raise ValueError(tr(f"Неверное время «{_str(v)}» (ожидается 10:00–10:50)", f"Invalid time “{_str(v)}” (expected 10:00–10:50)",
                            f"Уақыты қате: «{_str(v)}» (10:00–10:50 пішімі күтіледі)"))
    start, end = time(int(m[1]), int(m[2])), time(int(m[3]), int(m[4]))
    if end <= start:
        raise ValueError(tr("Время окончания раньше начала", "The end time is earlier than the start time",
                            "Аяқталу уақыты басталу уақытынан ерте"))
    return start, end


ROMAN = {"I": 1, "II": 2, "III": 3, "1": 1, "2": 2, "3": 3}


def find_location(db: Session, text: str) -> Location | None:
    m = re.search(r"C1\.\d\.\d{3}[A-ZА-Я]?", text.upper().replace(" ", ""))
    if m:
        code = m.group(0).replace("К", "K").replace("Р", "P")
        loc = db.scalar(select(Location).where(Location.code == code))
        if loc:
            return loc
    return db.scalar(select(Location).where(func.lower(Location.title) == text.lower()))


def import_schedule(db: Session, trimester: Trimester, content: bytes) -> dict:
    if trimester.status == TrimesterStatus.approved.value:
        raise HTTPException(400, tr("График утверждён — новые занятия добавляйте вручную",
                                    "The schedule is approved — add new lessons manually",
                                    "Кесте бекітілген — жаңа сабақтарды қолмен қосыңыз"))
    users = db.scalars(select(User)).all()
    by_name = {normalize_name(u.full_name): u for u in users}
    by_email = {u.email.lower(): u for u in users if u.email}
    created, errors = 0, []
    for n, row in _rows(content, SCHEDULE_FIELDS):
        try:
            year = _str(row.get("year")).replace("-", "–")
            selected_year = trimester.year.name
            if year and year != selected_year.replace("-", "–"):
                raise ValueError(tr(f"Учебный год {year} не совпадает с выбранным {selected_year}",
                                    f"Academic year {year} does not match the selected {selected_year}",
                                    f"{year} оқу жылы таңдалған {selected_year} оқу жылымен сәйкес келмейді"))
            tri_text = _str(row.get("trimester"))
            tri = ROMAN.get(tri_text.upper())
            if tri and tri != trimester.number:
                raise ValueError(tr(f"Триместр {tri_text} не совпадает с выбранным {trimester.roman}",
                                    f"Trimester {tri_text} does not match the selected {trimester.roman}",
                                    f"{tri_text} триместр таңдалған {trimester.roman} триместрмен сәйкес келмейді"))
            name = _str(row.get("teacher"))
            email = _str(row.get("email")).lower()
            teacher = by_email.get(email) if email else None
            teacher = teacher or by_name.get(normalize_name(name))
            if teacher is None:
                if not email:
                    raise ValueError(tr(f"Преподаватель «{name}» не найден в справочнике ППС (укажите Email)",
                                        f"Teacher “{name}” is not in the faculty directory (specify the e-mail)",
                                        f"Оқытушы «{name}» ПОҚ анықтамалығында табылмады (Email көрсетіңіз)"))
                teacher = User(email=email, full_name=name, roles=["teacher"])
                db.add(teacher)
                db.flush()
                by_email[email] = teacher
            position = _str(row.get("position")) or teacher.position
            if position and not teacher.position:
                teacher.position = position
            dept = _department(db, _str(row.get("department"))) or teacher.department
            if dept and not teacher.department_id:
                teacher.department_id = dept.id
            topic, discipline = _str(row.get("topic")), _str(row.get("discipline"))
            if not topic or not discipline:
                raise ValueError(tr("Не указаны тема или дисциплина", "The topic or the discipline is missing",
                                    "Тақырыбы немесе пәні көрсетілмеген"))
            d = _parse_date(row.get("date"))
            start, end = _parse_time_range(row.get("time"))
            place = _str(row.get("place"))
            loc = find_location(db, place) if place else None
            db.add(
                OpenLesson(
                    trimester_id=trimester.id,
                    teacher_id=teacher.id,
                    teacher_position=position,
                    department_id=dept.id if dept else None,
                    topic=topic,
                    discipline=discipline,
                    date=d,
                    start_time=start,
                    end_time=end,
                    location_id=loc.id if loc else None,
                    location_text=loc.code if loc else place,
                    status=LessonStatus.draft.value,
                    materials_deadline=lesson_logic.compute_deadline(db, d),
                )
            )
            db.flush()
            created += 1
        except ValueError as exc:
            errors.append({"row": n, "error": str(exc)})
    db.commit()
    return {"created": created, "errors": errors}
